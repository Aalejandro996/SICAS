from datetime import date, datetime, timezone
from pathlib import Path
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

from app.core.security import permissions as P
from app.modules.analytics.contralor_service import (
    METRICS, SCOPE_CTE, _excel_safe_text, _kpi_audit_snapshot, _kpi_status, _validate_kpi,
    export_workbook, parse_filters, period_window,
)


class ContralorDashboardUnitTests(unittest.TestCase):
    def test_private_dashboard_permissions_are_exclusive_to_contralor(self):
        self.assertIn("dashboard:contralor", P.ROLE_PERMISSIONS[P.CONTRALOR])
        self.assertIn("kpi:manage", P.ROLE_PERMISSIONS[P.CONTRALOR])
        for role, permissions in P.ROLE_PERMISSIONS.items():
            if role != P.CONTRALOR:
                self.assertNotIn("dashboard:contralor", permissions, role)
                self.assertNotIn("kpi:manage", permissions, role)
        self.assertIn("dashboard:executive", P.ROLE_PERMISSIONS[P.GERENCIA])
        self.assertNotIn("kpi:export", P.ROLE_PERMISSIONS[P.GERENCIA])

    def test_date_filters_validate_order_and_range(self):
        good = parse_filters({"start_date": "2026-01-01", "end_date": "2026-01-31", "granularity": "MONTH"})
        self.assertEqual(good["start_date"], date(2026, 1, 1))
        self.assertEqual(good["end_date"], date(2026, 1, 31))
        with self.assertRaises(ValueError):
            parse_filters({"start_date": "2026-02-01", "end_date": "2026-01-31"})
        with self.assertRaises(ValueError):
            parse_filters({"start_date": "2010-01-01", "end_date": "2026-10-01"})

    def test_period_window_returns_current_partial_period(self):
        self.assertEqual(period_window("MENSUAL", date(2026, 10, 13)), (date(2026, 10, 1), date(2026, 10, 13)))
        self.assertEqual(period_window("TRIMESTRAL", date(2026, 10, 13)), (date(2026, 10, 1), date(2026, 10, 13)))
        self.assertEqual(period_window("SEMESTRAL", date(2026, 10, 13)), (date(2026, 7, 1), date(2026, 10, 13)))
        self.assertEqual(period_window("ANUAL", date(2026, 10, 13)), (date(2026, 1, 1), date(2026, 10, 13)))

    def test_list_kpis_uses_historical_as_of_date(self):
        from app.modules.analytics import contralor_service as service

        row = {
            "id": "00000000-0000-0000-0000-000000000001",
            "name": "Auditorías mensuales",
            "description": "",
            "metric_code": "AUDITS_TOTAL",
            "target_value": 10,
            "warning_value": None,
            "direction": "AT_LEAST",
            "periodicity": "MENSUAL",
            "scope_type": "GENERAL",
            "unit_id": None,
            "auditor_id": None,
            "active": True,
            "created_at": None,
            "unit_name": "",
            "auditor_name": "",
        }
        as_of = date(2026, 9, 30)

        with (
            patch.object(service, "_mappings", return_value=[row]),
            patch.object(service, "metric_value", return_value=8.0) as metric,
        ):
            result = service.list_kpis(as_of=as_of)

        self.assertEqual(result[0]["period_start"], "2026-09-01")
        self.assertEqual(result[0]["period_end"], "2026-09-30")
        metric_filters = metric.call_args.args[1]
        self.assertEqual(metric_filters["start_date"], date(2026, 9, 1))
        self.assertEqual(metric_filters["end_date"], date(2026, 9, 30))

    def test_snapshot_anchors_kpis_to_filter_end_date(self):
        import inspect
        from app.modules.analytics import contralor_service as service

        source = inspect.getsource(service.create_dashboard_snapshot)
        self.assertIn(
            'kpis = list_kpis(as_of=filters["end_date"])',
            source,
        )

    def test_monthly_report_anchors_kpis_to_report_end_date(self):
        source = Path(
            "scripts/generate_contralor_monthly_report.py"
        ).read_text(encoding="utf-8")
        self.assertIn("kpis = service.list_kpis(as_of=end)", source)

    def test_kpi_metrics_are_allowlisted_and_scopes_are_validated(self):
        self.assertIn("AUDIT_COMPLETION_PCT", METRICS)
        base = {"name": "Cierre mensual", "metric_code": "AUDIT_COMPLETION_PCT", "target_value": 80,
                "direction": "AT_LEAST", "periodicity": "MENSUAL", "scope_type": "GENERAL"}
        self.assertEqual(_validate_kpi(base)["target_value"], 80.0)
        with self.assertRaises(ValueError):
            _validate_kpi({**base, "metric_code": "SELECT * FROM iam.users"})
        with self.assertRaises(ValueError):
            _validate_kpi({**base, "direction": "AT_LEAST", "warning_value": 90})
        with self.assertRaises(ValueError):
            _validate_kpi({**base, "metric_code": "INGESTION_AVG_SECONDS", "scope_type": "UNIDAD", "unit_id": 1})
        with self.assertRaises(ValueError):
            _validate_kpi({**base, "target_value": -1})
        with self.assertRaises(ValueError):
            _validate_kpi({**base, "metric_code": "AUDIT_COMPLETION_PCT", "target_value": 101})
        with self.assertRaises(ValueError):
            _validate_kpi({**base, "target_value": True})

    def test_kpi_unit_id_rejects_boolean_and_fractional_values(self):
        base = {
            "name": "Auditorías por unidad",
            "metric_code": "AUDITS_TOTAL",
            "target_value": 5,
            "direction": "AT_LEAST",
            "periodicity": "MENSUAL",
            "scope_type": "UNIDAD",
        }

        for value in (1, "1", 1.0):
            with self.subTest(valid_unit_id=value):
                result = _validate_kpi({**base, "unit_id": value})
                self.assertEqual(result["unit_id"], 1)

        for value in (1.5, True, False, 0, -1, "1.5", "abc",
                      float("nan"), float("inf")):
            with self.subTest(invalid_unit_id=value):
                with self.assertRaisesRegex(ValueError, "unidad elegida"):
                    _validate_kpi({**base, "unit_id": value})

    def test_closed_audit_duration_uses_pac_sent_timestamp(self):
        self.assertIn("a.pac_sent_at", SCOPE_CTE)
        self.assertIn("engagement.business_seconds(a.created_at, a.pac_sent_at) / 32400.0", SCOPE_CTE)   # días hábiles de 9 horas
        self.assertNotIn("lp.last_changed_at - a.created_at", SCOPE_CTE)

    def test_kpi_audit_snapshot_only_contains_configuration_fields(self):
        snapshot = _kpi_audit_snapshot({"name": "Cierre", "metric_code": "AUDITS_CLOSED", "active": True,
                                        "password": "no-log", "created_by": "actor"})
        self.assertEqual(snapshot, {"name": "Cierre", "metric_code": "AUDITS_CLOSED", "active": True})

    def test_export_workbook_contains_formulas_and_charts(self):
        data = {
            "filters": {"start_date": "2026-01-01", "end_date": "2026-10-13", "audit_domain": None, "unit_id": None},
            "summary": {"audits_total": 5, "audits_in_progress": 4, "audits_closed": 1, "pending_approval": 1,
                        "observations_total": 22, "observations_open": 22, "high_risk_observations": 20,
                        "avg_closed_audit_days": 4.3, "avg_open_days": 6.1},
            "by_unit": [{"unit_name": "DEMO", "unit_code": "1", "audits_total": 5, "audits_in_progress": 4,
                         "audits_closed": 1, "observations_total": 22, "high_risk_observations": 20,
                         "observations_open": 22, "avg_closed_audit_days": 4.3,
                         "risk_breakdown": {"CRITICAL": 0, "HIGH": 20, "MEDIUM": 2, "LOW": 0, "UNSPECIFIED": 0, "OTHER": 0}}],
            "by_domain": [{"audit_domain": "FINANCIERA", "audits_total": 1, "audits_in_progress": 1, "audits_closed": 0,
                           "avg_closed_audit_days": None, "observations_total": 0, "observations_open": 0,
                           "critical_observations": 0, "high_risk_observations": 0, "medium_observations": 0,
                           "low_observations": 0, "unclassified_observations": 0, "other_observations": 0,
                           "observations_without_item": 0, "criteria_with_formal_observation": 0,
                           "items_total": 102, "items_scored": 6, "items_scored_pct": 5.88, "high_risk_pct": None}],
            "by_auditor": [{"auditor_name": "Auditor Demo", "audits_total": 5, "audits_closed": 1, "avg_closed_audit_days": 4.3}],
            "observations_by_risk": {k: {"total": n, "open": n} for k, n in
                                     (("CRITICAL", 0), ("HIGH", 20), ("MEDIUM", 2), ("LOW", 0), ("UNSPECIFIED", 0), ("OTHER", 0))},
            "trend": [{"period_start": datetime(2026, 1, 1, tzinfo=timezone.utc), "audits_total": 5,
                       "audits_closed": 1, "observations_total": 22}],
            "phase_times": [{"phase": "INFORME", "transitions": 2, "average_hours": 8.3}],
            "ingestion": {"jobs_total": 2, "jobs_failed": 0, "jobs_in_progress": 0, "rows_total": 100,
                          "rows_processed": 95, "rows_rejected": 5, "average_seconds": 12.5, "rejection_pct": 5.0},
        }
        kpis = [{"name": '=HYPERLINK("https://invalid.example","Abrir")', "metric": {"label": "Cierre", "unit": "%"}, "periodicity": "MENSUAL",
                 "scope_type": "GENERAL", "actual_value": 75, "target_value": 80, "warning_value": 70,
                 "direction": "AT_LEAST", "period_start": "2026-10-01", "period_end": "2026-10-13"}]
        output = export_workbook(data, kpis)
        workbook = load_workbook(output, data_only=False)
        self.assertIn("Comparativo por unidad", workbook.sheetnames)
        self.assertIn("Comparativo por dominio", workbook.sheetnames)
        self.assertIn("Riesgos por unidad", workbook.sheetnames)
        self.assertEqual(workbook["Comparativo por dominio"]["R2"].value, '=IF(P2=0,"",Q2/P2)')
        self.assertEqual(workbook["Comparativo por dominio"]["S2"].value, '=IF(F2=0,"",(H2+I2)/F2)')
        self.assertEqual(len(workbook["Comparativo por dominio"]._charts), 1)
        self.assertIn("Comparativo por auditor", workbook.sheetnames)
        self.assertIn("Alertas actuales", workbook.sheetnames)
        self.assertIn("Calidad de datos", workbook.sheetnames)
        self.assertIn("Flujo y carga", workbook.sheetnames)
        self.assertIn("Cortes históricos", workbook.sheetnames)
        self.assertEqual(workbook["Comparativo por unidad"]["J2"].value, '=IF(F2=0,"",G2/F2)')
        self.assertTrue(workbook["KPI"]["I2"].value.startswith('=IF(E2='))
        self.assertTrue(workbook["KPI"]["A2"].value.startswith("'="))
        self.assertNotEqual(workbook["KPI"]["A2"].data_type, "f")
        self.assertEqual(len(workbook["Histórico"]._charts), 1)
        self.assertEqual(len(workbook["Comparativo por unidad"]._charts), 1)
        self.assertEqual(len(workbook["Riesgos por unidad"]._charts), 1)
        self.assertEqual(workbook["Riesgos por unidad"]["I2"].value, "=SUM(C2:H2)")

    def test_excel_text_neutralizes_formula_prefixes(self):
        for value in ('=1+1', '  +1+1', '-1+1', '@SUM(A1:A2)'):
            self.assertTrue(_excel_safe_text(value).startswith("'"))
        self.assertEqual(_excel_safe_text("Texto normal"), "Texto normal")

    def test_kpi_evaluation_handles_missing_data_and_warning_band(self):
        self.assertEqual(_kpi_status(None, 80, 70, "AT_LEAST"), "SIN_DATOS")
        self.assertEqual(_kpi_status(82, 80, 70, "AT_LEAST"), "CUMPLIDO")
        self.assertEqual(_kpi_status(75, 80, 70, "AT_LEAST"), "ATENCION")
        self.assertEqual(_kpi_status(60, 80, 70, "AT_LEAST"), "DESVIACION")
        self.assertEqual(_kpi_status(4, 5, 7, "AT_MOST"), "CUMPLIDO")
        self.assertEqual(_kpi_status(6, 5, 7, "AT_MOST"), "ATENCION")


    def test_alert_severity_and_summary_are_documented_in_service(self):
        import inspect
        from app.modules.analytics import contralor_service as service
        src = inspect.getsource(service.get_control_alerts)
        self.assertIn("AUDITORIA_VENCIDA", src)
        self.assertIn("PAC_VENCIDO", src)
        self.assertIn("BRECHA_ALTO_CRITICA", src)

    def test_quality_checks_treat_general_observations_as_info(self):
        import inspect
        from app.modules.analytics import contralor_service as service
        src = inspect.getsource(service.get_data_quality)
        self.assertIn('"GENERAL_OBSERVATIONS"', src)
        self.assertIn('"INFO"', src)
        self.assertIn('"CROSS_AUDIT_OBSERVATION_LINK"', src)

    def test_snapshot_source_allowlist_and_html_escaping_are_present(self):
        import inspect
        from app.modules.analytics import contralor_service as service
        src = inspect.getsource(service.create_dashboard_snapshot)
        self.assertIn('"MANUAL", "AUTOMATICO"', src)
        self.assertIn("html.escape", inspect.getsource(service.render_print_report))

    def test_print_report_contains_save_to_pdf_action(self):
        from app.modules.analytics.contralor_service import render_print_report
        data = {"filters": {"start_date":"2026-10-01", "end_date":"2026-10-31", "audit_domain":None, "unit_id":None, "auditor_id":None},
                "summary":{"audits_total":0,"audits_in_progress":0,"audits_closed":0,"observations_open":0,"high_risk_observations":0},
                "alert_summary":{"total":0}, "data_quality":{"issues_total":0,"checks":[]}, "by_domain":[], "alerts":[],
                "phase_pipeline":{"phases":[]}, "workload":[]}
        report = render_print_report(data, [])
        self.assertIn("Imprimir / Guardar como PDF", report)
        self.assertIn("Comparación FINANCIERA vs SISTEMAS", report)

    def test_pdf_export_produces_valid_pdf_header(self):
        from app.modules.analytics.contralor_service import export_pdf_report
        data = {"filters": {"start_date":"2026-10-01", "end_date":"2026-10-31", "audit_domain":None, "unit_id":None, "auditor_id":None},
                "summary":{"audits_total":0,"audits_in_progress":0,"audits_closed":0,"observations_open":0,"high_risk_observations":0},
                "alert_summary":{"total":0}, "data_quality":{"issues_total":0,"checks":[]}, "by_domain":[], "alerts":[],
                "phase_pipeline":{"phases":[]}, "workload":[]}
        pdf = export_pdf_report(data, [])
        self.assertTrue(pdf.getvalue().startswith(b"%PDF-"))



    def test_activity_sessions_are_explicit_and_one_active_session_is_enforced(self):
        from pathlib import Path
        migration = Path("sql/016_auditor_activity_tracking.sql").read_text(encoding="utf-8")
        self.assertIn("uq_auditor_activity_one_open_session", migration)
        self.assertIn("Las sesiones finalizadas son inmutables", migration)
        import inspect
        from app.modules.analytics import contralor_service as service
        self.assertIn("start_activity_session", inspect.getsource(service))
        self.assertIn("last_resumed_at", inspect.getsource(service.transition_activity_session))

    def test_activity_duration_formatter_handles_pause_accumulation(self):
        from pathlib import Path
        page = Path("app/static/pages_activity.js").read_text(encoding="utf-8")
        self.assertIn("Pausar", page)
        self.assertIn("Reanudar", page)
        self.assertIn("Finalizar sesión", page)
        self.assertIn("setInterval", page)


    def test_completed_pac_is_not_counted_as_open_breach_and_financial_counts_use_snapshot(self):
        import inspect
        from app.modules.analytics import contralor_service as service
        self.assertIn("'CUMPLIDA'", service.OBS_CTE)
        self.assertIn("'CUMPLIDO'", service.OBS_CTE)
        domain_source = inspect.getsource(service._get_domain_comparison)
        self.assertIn("'CUMPLIDO'", domain_source)
        alerts_source = inspect.getsource(service.get_control_alerts)
        self.assertIn("NOT IN ('CUMPLIDA','CUMPLIDO'", alerts_source)
        workload_source = inspect.getsource(service.get_auditor_workload)
        self.assertIn("NOT IN ('CUMPLIDA','CUMPLIDO'", workload_source)
        quality_source = inspect.getsource(service.get_data_quality)
        self.assertIn("engagement.financial_worksheets fw", quality_source)
        self.assertIn("fw.criteria_snapshot", quality_source)
        workload_source = inspect.getsource(service.get_auditor_workload)
        self.assertIn("AS elapsed_seconds", inspect.getsource(service.get_activity_options))
        self.assertIn("last_resumed_at", workload_source)


if __name__ == "__main__":
    unittest.main()
