import io
import unittest
from datetime import datetime, timezone

from openpyxl import load_workbook

from app.modules.engagement import offline, scoring
from app.modules.engagement.pac import build_pac
from app.modules.engagement.templates import DEFAULTS, default_criteria, validate_criteria, branches_for_line, KINDS, LINES

SECRET = b"k" * 32
NOW = datetime(2026, 10, 2, 9, 0, tzinfo=timezone.utc)


class TemplatesTest(unittest.TestCase):
    def test_six_templates_exist_and_are_valid(self):
        self.assertEqual(set(DEFAULTS), {(k, l) for k in KINDS for l in LINES})        # 2 tipos x 3 líneas, separadas
        for (k, l), crit in DEFAULTS.items():
            self.assertEqual(validate_criteria(crit), crit)
            self.assertGreaterEqual(len(crit), 9)

    def test_validation_rejects_bad_input(self):
        for bad in ([], [{"code": "A", "text": "texto largo", "weight": 1}], [{"code": "S01", "text": "corto", "weight": 5}],
                    [{"code": "S01", "text": "texto válido", "weight": 1}, {"code": "s01", "text": "otro válido", "weight": 1}]):
            with self.assertRaises(ValueError):
                validate_criteria(bad)

    def test_branch_mapping(self):
        self.assertIn("AUTO", branches_for_line("AUTO")); self.assertIn("HOGAR", branches_for_line("PATRIMONIAL"))
        self.assertIn("VIDA", branches_for_line("PERSONAL"))


class ScoringTest(unittest.TestCase):
    CR = [{"code": "A1", "weight": 3}, {"code": "A2", "weight": 1}, {"code": "A3", "weight": 2}]

    def test_weighted_score_and_excluded_results(self):
        self.assertEqual(scoring.item_score(self.CR, {"A1": "CUMPLE", "A2": "NO_CUMPLE", "A3": "PARCIAL"}), round((300 + 0 + 100) / 6, 2))
        self.assertEqual(scoring.item_score(self.CR, {"A1": "CUMPLE", "A2": "NO_APLICA", "A3": "EXCEPCION"}), 100.0)  # excluidos
        self.assertIsNone(scoring.item_score(self.CR, {"A1": "EXCEPCION"}))
        self.assertIsNone(scoring.item_score(self.CR, {}))

    def test_manual_override_requires_reason_and_range(self):
        self.assertEqual(scoring.validate_manual("72.5", "Subsanado durante la visita"), (72.5, "Subsanado durante la visita"))
        for bad in (("101", "motivo suficiente largo"), ("-1", "motivo suficiente largo"), ("abc", "motivo suficiente"), ("50", "corto")):
            with self.assertRaises(scoring.ScoreError):
                scoring.validate_manual(*bad)
        self.assertEqual(scoring.final_score(40.0, 75.0), 75.0)
        self.assertEqual(scoring.final_score(40.0, None), 40.0)
        self.assertEqual(scoring.final_score(40.0, 0.0), 0.0)         # un 0 manual es válido (no se confunde con "sin ajuste")

    def test_decimals_from_database_are_handled(self):
        from decimal import Decimal
        items = [{"results": {"A1": "CUMPLE"}, "manual_score": Decimal("60.00")}, {"results": {"A1": "CUMPLE"}, "manual_score": None}]
        s = scoring.worksheet_summary(self.CR, items, sheet_manual=Decimal("90.00"))
        self.assertEqual((s.auto, s.final), (80.0, 90.0))
        self.assertEqual(scoring.audit_summary([s], manual=Decimal("70"))["final"], 70.0)
        self.assertIsInstance(scoring.final_score(10.0, Decimal("55.5")), float)

    def test_rating_thresholds(self):
        self.assertEqual(scoring.rating(90), "SATISFACTORIO"); self.assertEqual(scoring.rating(85), "SATISFACTORIO")
        self.assertEqual(scoring.rating(84.99), "NECESITA_MEJORAS"); self.assertEqual(scoring.rating(70), "NECESITA_MEJORAS")
        self.assertEqual(scoring.rating(69.9), "INSATISFACTORIO"); self.assertEqual(scoring.rating(None), "SIN_CALIFICAR")
        self.assertEqual(scoring.rating(80, {"satisfactorio": 80}), "SATISFACTORIO")

    def test_worksheet_and_audit_summary_with_overrides_and_exceptions(self):
        items = [{"results": {"A1": "CUMPLE", "A2": "CUMPLE", "A3": "CUMPLE"}},
                 {"results": {"A1": "NO_CUMPLE", "A2": "NO_CUMPLE", "A3": "NO_CUMPLE"}, "manual_score": 60},
                 {"results": {"A1": "EXCEPCION"}, "info_exception": "No entregaron el expediente"}]
        s = scoring.worksheet_summary(self.CR, items)
        self.assertEqual((s.items, s.scored, s.exceptions, s.auto), (3, 2, 1, 80.0))        # (100 + 60) / 2
        self.assertEqual(s.rating, "NECESITA_MEJORAS")
        s2 = scoring.worksheet_summary(self.CR, items, sheet_manual=95)
        self.assertTrue(s2.manual_override); self.assertEqual((s2.final, s2.auto), (95, 80.0)); self.assertEqual(s2.rating, "SATISFACTORIO")
        a = scoring.audit_summary([s, s2])
        self.assertEqual(a["auto"], 87.5)
        self.assertEqual(scoring.audit_summary([s, s2], manual=50)["rating"], "INSATISFACTORIO")
        self.assertEqual(a["exceptions"], 2)


def make_sheets():
    crit = default_criteria("RECLAMO", "AUTO")
    codes = [c["code"] for c in crit]
    items = [{"id": f"id-{i}", "ref": f"SIN-{i:04d}", "results": {codes[0]: "CUMPLE"}, "manual_score": None, "manual_reason": None,
              "observation": "=HYPERLINK(\"http://x\")" if i == 1 else None, "info_exception": None, "updated_at": NOW} for i in range(1, 4)]
    return [{"id": "ws-1", "kind": "RECLAMO", "line": "AUTO", "title": "t", "criteria": crit, "items": items}], codes


class OfflineTest(unittest.TestCase):
    def setUp(self):
        self.sheets, self.codes = make_sheets()
        self.audit = {"id": "aud-1", "code": "AUD-2026-0001", "title": "Auditoría de prueba", "unit": "Sucursal Centro"}
        self.data = offline.export_workbook(self.audit, self.sheets, SECRET, NOW, "auditor1")

    def edit(self, fn):
        wb = load_workbook(io.BytesIO(self.data)); fn(wb["REC-AUTO"], wb)
        out = io.BytesIO(); wb.save(out); return out.getvalue()

    def current(self, changed=()):
        cur = {}
        for it in self.sheets[0]["items"]:
            obs = "CAMBIO EN SERVIDOR" if it["id"] in changed else it["observation"]
            cur[it["id"]] = {"hash": offline.row_hash(it["results"], None, None, obs, None, self.codes)}
        return cur

    def run_import(self, data, **kw):
        control, rows, errs = offline.parse_workbook(data, SECRET, "aud-1")
        refs = {"ws-1": {i["ref"] for i in self.sheets[0]["items"]}}
        return control, rows, errs, offline.plan_changes(control, rows, kw.get("current", self.current()), refs, kw.get("overwrite", False))

    def test_workbook_structure(self):
        wb = load_workbook(io.BytesIO(self.data))
        self.assertEqual(wb.sheetnames, ["Instrucciones", "Criterios", "REC-AUTO", "_control"])
        self.assertEqual(wb["_control"].sheet_state, "hidden")
        ws = wb["REC-AUTO"]
        self.assertEqual(ws.cell(3, 3).value, self.codes[0])
        self.assertEqual(ws.cell(4, 2).value, "SIN-0001")

    def test_formula_injection_is_neutralised(self):
        wb = load_workbook(io.BytesIO(self.data))
        c = wb["REC-AUTO"].cell(4, 2 + len(self.codes) + 3)
        self.assertEqual(c.data_type, "s")
        self.assertTrue(str(c.value).startswith("="))

    def test_unchanged_roundtrip(self):
        _, rows, errs, plan = self.run_import(self.data)
        self.assertEqual(errs, [])
        self.assertEqual(offline.summarize(plan)["UNCHANGED"], 3)

    def test_update_new_and_invalid_rows(self):
        c0 = 3
        def f(ws, wb):
            ws.cell(5, c0 + 1).value = "NO_CUMPLE"                        # actualiza SIN-0002
            ws.cell(5, c0 + len(self.codes)).value = 55                    # manual
            ws.cell(5, c0 + len(self.codes) + 1).value = "Subsanado durante la visita"
            ws.cell(7, 2).value = "SIN-9999"; ws.cell(7, c0).value = "PARCIAL"                  # fila nueva
            ws.cell(8, 2).value = "SIN-8888"; ws.cell(8, c0).value = "TAL VEZ"                  # valor inválido
            ws.cell(9, 2).value = "SIN-7777"; ws.cell(9, c0).value = "EXCEPCION"                # excepción sin detalle
        _, rows, errs, plan = self.run_import(self.edit(f))
        st = {p["ref"]: p["status"] for p in plan}
        self.assertEqual(st["SIN-0002"], "UPDATE"); self.assertEqual(st["SIN-9999"], "NEW")
        self.assertEqual(st["SIN-8888"], "INVALID"); self.assertEqual(st["SIN-7777"], "INVALID")
        upd = next(p for p in plan if p["ref"] == "SIN-0002")["data"]
        self.assertEqual((upd["manual"], upd["results"][self.codes[1]]), (55.0, "NO_CUMPLE"))

    def test_manual_score_needs_reason(self):
        def f(ws, wb): ws.cell(4, 3 + len(self.codes)).value = 50
        *_, plan = self.run_import(self.edit(f))
        self.assertEqual(plan[0]["status"], "INVALID")

    def test_conflict_when_server_changed_and_overwrite(self):
        def f(ws, wb): ws.cell(5, 3 + len(self.codes) + 2).value = "Observación hecha en campo"
        data = self.edit(f)
        *_, plan = self.run_import(data, current=self.current(changed={"id-2"}))
        self.assertEqual(next(p for p in plan if p["item_id"] == "id-2")["status"], "CONFLICT")
        *_, plan2 = self.run_import(data, current=self.current(changed={"id-2"}), overwrite=True)
        self.assertEqual(next(p for p in plan2 if p["item_id"] == "id-2")["status"], "UPDATE")

    def test_new_row_duplicates(self):
        # (a) misma referencia dos veces dentro de la hoja descargada -> inválida
        def f(ws, wb): ws.cell(7, 2).value = "sin-0001"; ws.cell(7, 3).value = "CUMPLE"
        *_, plan = self.run_import(self.edit(f))
        self.assertEqual(next(p for p in plan if p["row"] == 7)["status"], "INVALID")
        # (b) otra persona creó esa referencia en el servidor después de la descarga -> conflicto
        def g(ws, wb): ws.cell(7, 2).value = "SIN-5555"; ws.cell(7, 3).value = "CUMPLE"
        control, rows, _ = offline.parse_workbook(self.edit(g), SECRET, "aud-1")
        refs = {"ws-1": {i["ref"] for i in self.sheets[0]["items"]} | {"SIN-5555"}}
        plan = offline.plan_changes(control, rows, self.current(), refs)
        self.assertEqual(next(p for p in plan if p["row"] == 7)["status"], "CONFLICT")

    def test_tampered_control_or_foreign_file_rejected(self):
        wb = load_workbook(io.BytesIO(self.data)); wb["_control"]["A3"].value = wb["_control"]["A3"].value.replace("id-1", "id-X")
        out = io.BytesIO(); wb.save(out)
        with self.assertRaises(offline.OfflineError): offline.parse_workbook(out.getvalue(), SECRET, "aud-1")
        with self.assertRaises(offline.OfflineError): offline.parse_workbook(self.data, b"otra-clave" * 4, "aud-1")
        with self.assertRaises(offline.OfflineError): offline.parse_workbook(self.data, SECRET, "otra-auditoria")
        with self.assertRaises(offline.OfflineError): offline.parse_workbook(b"no es excel", SECRET, "aud-1")
        plain = io.BytesIO(); load_workbook(io.BytesIO(self.data)).save(plain)
        wb2 = load_workbook(io.BytesIO(self.data)); del wb2["_control"]; o2 = io.BytesIO(); wb2.save(o2)
        with self.assertRaises(offline.OfflineError): offline.parse_workbook(o2.getvalue(), SECRET, "aud-1")

    def test_large_control_is_chunked(self):
        crit = default_criteria("SUSCRIPCION", "PERSONAL")
        items = [{"id": f"00000000-0000-0000-0000-{i:012d}", "ref": f"POL-{i}", "results": {}, "updated_at": NOW} for i in range(1500)]
        big = offline.export_workbook(self.audit, [{"id": "ws-9", "kind": "SUSCRIPCION", "line": "PERSONAL", "title": "x", "criteria": crit, "items": items}], SECRET, NOW, "a")
        control, rows, errs = offline.parse_workbook(big, SECRET, "aud-1")
        self.assertEqual(len(control["sheets"]["SUS-PERSO"]["items"]), 1500); self.assertEqual(len(rows), 1500)


class PacTest(unittest.TestCase):
    def test_pac_workbook(self):
        obs = [{"seq": 1, "title": "Falta peritaje", "risk": "ALTO", "recommendation": "Exigirlo", "unit_decision": "ACEPTA",
                "pac_action": "=1+1", "pac_owner": "Gerente", "pac_due": None, "pac_status": "EN_CURSO"}]
        data = build_pac({"code": "AUD-1", "unit": "Sucursal"}, obs, NOW)
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["PAC", "Instrucciones"])
        ws = wb["PAC"]
        self.assertEqual((ws["B6"].value, ws["C6"].value, ws["I6"].value), ("Falta peritaje", "ALTO", "EN_CURSO"))
        self.assertEqual(ws["F6"].data_type, "s")                        # sin fórmulas inyectadas


if __name__ == "__main__":
    unittest.main()
