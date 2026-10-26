"""Importación segura del cuestionario oficial de Auditoría de Sistemas / ERPs TI."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from .templates import SYSTEM_ERP_CONTROLS

MAX_BYTES = 25 * 1024 * 1024
SHEET_NAME = "ERPs TI"
SECURITY_SHEET_NAME = "Seguridad_de_TI"


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().upper().replace("Á", "A").replace("É", "E").replace("Í", "I").replace("Ó", "O").replace("Ú", "U").split())


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _control(value: Any) -> str:
    v = _norm(value)
    if v in {"SI", "SÍ"}:
        return "SI"
    if v in {"NO"}:
        return "NO"
    # PARCIAL no es una respuesta válida en la fórmula de Sistemas (solo SI/NO):
    # el control queda pendiente para que el auditor lo responda en SICAS.
    return ""


def _doc(value: Any) -> str:
    v = _norm(value)
    return "SI" if v in {"SI", "SÍ"} else "NO" if v == "NO" else ""


def _evidence(value: Any) -> str:
    v = _norm(value)
    if v == "ACTUALIZADO":
        return "ACTUALIZADO"
    if v == "DESACTUALIZADO":
        return "DESACTUALIZADO"
    if v == "INEXISTENTE":
        return "INEXISTENTE"
    return ""

def _parse_generic_system_questionnaire(wb, filename: str | None = None) -> dict:
    """
    Lee un cuestionario Excel de Auditoría de Sistemas por su estructura.

    No depende de:
      - nombre del archivo;
      - nombre de la hoja;
      - ERP_TI;
      - Seguridad_de_TI;
      - Servidores;
      - Base de Datos;
      - cantidad fija de controles.

    Detecta una tabla que contenga como mínimo:
      ID
      Pregunta de Auditoría
      Existencia del Control
      Soporte Documental
      Evidencia / Entregable

    Conserva el ID original del Excel para trazabilidad.
    """

    import math
    import re
    import unicodedata

    def text(value) -> str:
        if value is None:
            return ""

        if isinstance(value, float) and math.isnan(value):
            return ""

        return str(value).replace("\u00a0", " ").strip()

    def header_key(value) -> str:
        """
        Normalización fuerte para encabezados.

        Permite reconocer, por ejemplo:
          Pregunta de Auditoría
          PreguntadeAuditoría
          PREGUNTA DE AUDITORIA
        """
        value = text(value)

        value = unicodedata.normalize("NFKD", value)

        value = "".join(
            ch for ch in value
            if not unicodedata.combining(ch)
        )

        return re.sub(
            r"[^a-z0-9]+",
            "",
            value.lower(),
        )

    HEADER_ALIASES = {
        "id": {
            "id",
            "idcontrol",
            "codigo",
            "codigocontrol",
            "codigodelcontrol",
            "controlid",
            "iddelcontrol",
        },
        "question": {
            "pregunta",
            "preguntadeauditoria",
            "preguntaauditoria",
            "preguntadelcontrol",
            "preguntadecontrol",
        },
        "control": {
            "existenciadelcontrol",
            "existenciacontrol",
            "existencia",
            "existe",
            "respuesta",
            "cumplimiento",
        },
        "documental": {
            "soportedocumental",
            "soporte",
            "documental",
        },
        "evidence": {
            "evidenciaentregable",
            "evidencia",
            "entregable",
        },
        "score": {
            "puntuacion",
            "puntuacion%",
            "puntuacionporcentual",
            "puntaje",
            "calificacion",
            "score",
        },
        "observation": {
            "observaciones",
            "observacion",
            "detalle",
            "comentarios",
            "comentario",
            "justificacion",
            "justificaciones",
        },
    }

    def detect_header(row):
        positions = {}

        for index, cell in enumerate(row):
            key = header_key(cell.value)

            if not key:
                continue

            for field, aliases in HEADER_ALIASES.items():
                if key in aliases and field not in positions:
                    positions[field] = index
                    break

        required = {
            "id",
            "question",
            "control",
            "documental",
            "evidence",
        }

        if required.issubset(positions):
            return positions

        return None

    def control_id(cell) -> str:
        """
        Conserva el ID tal como aparece en Excel.

        Ejemplos soportados:
          1
          01
          1.1
          01.01
          1.1.1
          GOB-01
          ACC-01
          AC-A2-37-08-2026-3
        """
        value = cell.value

        if value is None:
            return ""

        if isinstance(value, float) and math.isnan(value):
            return ""

        number_format = getattr(cell, "number_format", "") or ""

        if isinstance(value, float) and value.is_integer():
            value = int(value)

        if isinstance(value, int):
            fmt = number_format.split(";")[0].strip()

            if re.fullmatch(r"0+", fmt):
                return f"{value:0{len(fmt)}d}"

            return str(value)

        return text(value)

    def valid_control_id(value) -> bool:
        value = text(value)

        if not value:
            return False

        key = header_key(value)

        if key in {
            "id",
            "idcontrol",
            "codigo",
            "codigocontrol",
            "controlid",
        }:
            return False

        # Numéricos:
        # 01
        # 1
        # 1.1
        # 01.01
        # 1.1.1
        #
        # Alfanuméricos:
        # GOB-01
        # ACC-01
        # AC-A2-37-08-2026-3
        return bool(
            re.fullmatch(
                r"(?:\d+(?:[._-]\d+)*)|"
                r"(?:[A-Za-z0-9]+(?:[._-][A-Za-z0-9]+)+)",
                value,
            )
        )

    def looks_like_section(value) -> bool:
        value = text(value)

        if not value:
            return False

        return bool(
            re.match(
                r"^\d+(?:\.\d+)*[.)]?\s+.+",
                value,
            )
        )

    def get_section(rows, header_index, default):
        """
        Busca el título de sección inmediatamente anterior
        al encabezado de la tabla.
        """
        for row in reversed(rows[:header_index]):
            values = [
                text(cell.value)
                for cell in row[:10]
                if text(cell.value)
            ]

            if len(values) != 1:
                continue

            candidate = values[0]

            if looks_like_section(candidate):
                return candidate

        return default

    def observation_value(row, index):
        """
        En los cuestionarios reales, "Observaciones" suele estar
        en H y el contenido real puede estar en I.

        Se recorren algunas celdas posteriores para capturar el
        comentario real y descartar placeholders.
        """
        if index is None:
            return None

        placeholders = {
            "detalle:",
            "(indique la justificacion, evidencia o plan de accion si aplica...)",
            "(indique nombre del responsable...)",
        }

        values = []

        for pos in range(
            index,
            min(index + 3, len(row)),
        ):
            value = text(row[pos].value)

            if not value:
                continue

            if header_key(value) in {
                header_key(x)
                for x in placeholders
            }:
                continue

            values.append(value)

        if not values:
            return None

        return " | ".join(
            dict.fromkeys(values)
        )[:2000] or None

    rows_found = []
    seen = {}
    sheets = []

    for ws in wb.worksheets:

        rows = list(
            ws.iter_rows(
                min_row=1,
                max_row=ws.max_row,
                max_col=min(ws.max_column, 30),
            )
        )

        if not rows:
            continue

        header = None
        header_index = None

        # Buscar la primera tabla de controles de la hoja.
        for idx, row in enumerate(rows):
            detected = detect_header(row)

            if detected:
                header = detected
                header_index = idx
                break

        if header is None:
            continue

        section = get_section(
            rows,
            header_index,
            ws.title,
        )

        count = 0

        for row_number, row in enumerate(
            rows[header_index + 1:],
            start=header_index + 2,
        ):

            # Algunos libros vuelven a colocar encabezados entre secciones.
            detected = detect_header(row)

            if detected:
                header = detected
                continue

            id_cell = row[header["id"]]

            current_id = control_id(id_cell)

            question = text(
                row[header["question"]].value
            )

            # Fila que no corresponde a un control.
            if not valid_control_id(current_id) or not question:

                candidate = ""

                for cell in row:
                    value = text(cell.value)

                    if value:
                        candidate = value
                        break

                if looks_like_section(candidate):
                    section = candidate

                continue

            unique_key = current_id.upper()

            if unique_key in seen:
                raise ValueError(
                    f"El control {current_id} está repetido en el Excel "
                    f"({seen[unique_key]} y {ws.title}:{row_number})"
                )

            seen[unique_key] = (
                f"{ws.title}:{row_number}"
            )

            score_value = None

            if header.get("score") is not None:
                score_value = row[
                    header["score"]
                ].value

            rows_found.append(
                {
                    "control_id": current_id,
                    "source_id": current_id,
                    "ref": current_id,
                    "question": question,
                    "section": section or ws.title,
                    "control": text(
                        row[header["control"]].value
                    ),
                    "documental": text(
                        row[header["documental"]].value
                    ),
                    "evidencia": text(
                        row[header["evidence"]].value
                    ),
                    "observation": observation_value(
                        row,
                        header.get("observation"),
                    ),
                    "score": score_value,
                    "source_sheet": ws.title,
                    "source_row": row_number,
                }
            )

            count += 1

        if count:
            sheets.append(
                {
                    "name": ws.title,
                    "controls": count,
                }
            )

    if not rows_found:
        raise ValueError(
            "No se encontró una tabla de controles válida. "
            "El Excel debe contener columnas equivalentes a "
            "ID, Pregunta de Auditoría, Existencia del Control, "
            "Soporte Documental y Evidencia / Entregable."
        )

    return {
        "questionnaire_type": "SYSTEM_GENERIC",
        "questionnaire_name": (
            text(filename)
            if filename
            else "Cuestionario de Sistemas"
        )[:200],
        "source_filename": (
            text(filename)[:255]
            if filename
            else None
        ),
        "sheets": sheets,
        "rows": rows_found,
        "controls": len(rows_found),
    }


def parse(data: bytes, filename: str | None = None) -> dict:
    """
    Entrada única para cualquier cuestionario Excel de Sistemas.

    NO exige:
      - hoja ERPs TI;
      - hoja Seguridad_de_TI;
      - hoja Servidores;
      - nombre específico;
      - cantidad específica de controles.
    """
    if not data:
        raise ValueError(
            "El archivo está vacío"
        )

    if len(data) > MAX_BYTES:
        raise ValueError(
            "El cuestionario supera 25 MB"
        )

    if not data.startswith(b"PK"):
        raise ValueError(
            "El archivo no es un XLSX válido"
        )

    try:
        wb = load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:  # noqa: BLE001
        raise ValueError(
            "No fue posible leer el archivo Excel"
        ) from exc

    try:
        return _parse_generic_system_questionnaire(
            wb,
            filename=filename,
        )
    finally:
        try:
            wb.close()
        except Exception:
            pass