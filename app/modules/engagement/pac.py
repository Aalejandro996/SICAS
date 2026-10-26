"""PAC — Plan de Acción Correctiva (anexo en Excel que se envía a la unidad/sucursal)."""
from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ARIAL = "Arial"
COLS = [("N°", 6), ("Observación", 42), ("Riesgo", 9), ("Recomendación de Auditoría", 46), ("Respuesta de la unidad", 16),
        ("Acción correctiva propuesta", 44), ("Responsable", 24), ("Fecha compromiso", 16), ("Estatus", 14)]


def _s(ws, r, c, v):
    cell = ws.cell(row=r, column=c)
    cell.value = v
    if isinstance(v, str) and v[:1] in "=+-@":
        cell.data_type = "s"
    return cell


def build_pac(audit: dict, observations: list[dict], generated_at: datetime) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "PAC"
    navy = PatternFill("solid", fgColor="1F3864")
    fill_in = PatternFill("solid", fgColor="FFF9E5")
    thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "PLAN DE ACCIÓN CORRECTIVA (PAC)"; ws["A1"].font = Font(name=ARIAL, bold=True, size=14, color="1F3864")
    ws["A2"] = f"Auditoría {audit['code']} — {audit['unit']} · Emitido el {generated_at:%d/%m/%Y}"; ws["A2"].font = Font(name=ARIAL, size=10)
    ws["A3"] = "La unidad completa las columnas amarillas y devuelve este archivo a la Contraloría."; ws["A3"].font = Font(name=ARIAL, italic=True, size=9)
    hdr = 5
    for j, (h, w) in enumerate(COLS, 1):
        c = ws.cell(row=hdr, column=j, value=h)
        c.font, c.fill, c.border = Font(name=ARIAL, bold=True, color="FFFFFF", size=10), navy, box
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[chr(64 + j)].width = w
    ws.row_dimensions[hdr].height = 32
    for i, o in enumerate(observations, 1):
        r = hdr + i
        vals = [i, o["title"], o["risk"], o.get("recommendation") or "", o.get("unit_decision") or "",
                o.get("pac_action") or "", o.get("pac_owner") or "", o.get("pac_due"), o.get("pac_status") or "PENDIENTE"]
        for j, v in enumerate(vals, 1):
            c = _s(ws, r, j, v)
            c.font, c.border = Font(name=ARIAL, size=10), box
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j in (5, 6, 7, 8, 9):
                c.fill = fill_in
        if isinstance(vals[7], (date, datetime)):
            ws.cell(row=r, column=8).number_format = "DD/MM/YYYY"
    last = hdr + max(len(observations), 1) + 20
    dv1 = DataValidation(type="list", formula1='"ACEPTA,NO_ACEPTA,PARCIAL"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"PENDIENTE,EN_CURSO,CUMPLIDA"', allow_blank=True)
    ws.add_data_validation(dv1); ws.add_data_validation(dv2)
    dv1.add(f"E{hdr + 1}:E{last}"); dv2.add(f"I{hdr + 1}:I{last}")
    ws.freeze_panes = f"C{hdr + 1}"
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True

    ins = wb.create_sheet("Instrucciones")
    for i, t in enumerate([
        "Cómo completar el PAC", "",
        "• Respuesta de la unidad: ACEPTA, NO_ACEPTA o PARCIAL (si no acepta, explique en la acción propuesta).",
        "• Acción correctiva: describa qué se hará para corregir la causa, no solo el síntoma.",
        "• Responsable: nombre y cargo. Fecha compromiso: fecha realista de cumplimiento (DD/MM/AAAA).",
        "• Estatus: PENDIENTE, EN_CURSO o CUMPLIDA.", "",
        "Ejemplo:  1 | Falta peritaje firmado | ALTO | Exigir peritaje antes de aprobar | ACEPTA | "
        "Checklist obligatorio en el sistema de reclamos | Gerente de Siniestros | 30/11/2026 | EN_CURSO"], 1):
        c = ins.cell(row=i, column=1, value=t); c.font = Font(name=ARIAL, bold=(i == 1), size=12 if i == 1 else 10)
    ins.column_dimensions["A"].width = 120
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()
