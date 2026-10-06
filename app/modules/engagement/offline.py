"""
1. Trabajo fuera de la oficina: paquete de trabajo en Excel (.xlsx) que se descarga, se llena sin conexión y se carga al volver.

Seguridad e integridad:
  - La hoja oculta "_control" lleva los identificadores de cédulas/ítems, la huella de lo exportado y un HMAC-SHA256
    firmado con una clave del servidor: un archivo alterado o ajeno a SICAS se rechaza.
  - Detección de conflictos: si el ítem cambió en el servidor DESPUÉS de la descarga, la fila no se pisa (salvo que el
    Auditor lo ordene explícitamente tras ver la vista previa).
  - Todo texto que empiece por = + - @ se escribe como texto (evita inyección de fórmulas al abrir en Excel).
"""
from __future__ import annotations

import hashlib
import hmac
import io
import json
import zipfile
from datetime import datetime

def _xlsx():
    """Carga openpyxl de forma diferida para que SICAS pueda arrancar aunque el
    componente Excel no esté instalado; las operaciones offline devuelven un
    error controlado en lugar de provocar un 500 durante el arranque."""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        return Workbook, load_workbook, Alignment, Border, Font, PatternFill, Side, get_column_letter, DataValidation
    except ModuleNotFoundError as exc:
        raise OfflineError("El componente Excel no está instalado en el servidor. Instale openpyxl y reconstruya el contenedor API.") from exc

from .templates import KINDS, LINES, RESULTS, RESULT_LABEL

MAGIC = "SICAS-OFFLINE-1"
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
CHUNK = 30000
HEADER_ROW, FIRST_ROW = 3, 4
FIXED_LEFT = 2                       # ID, Referencia
TAIL = ["Calificación manual (0-100)", "Justificación de la calificación manual", "Observación del auditor",
        "Excepción: información no recibida (detalle)"]
ARIAL = "Arial"


class OfflineError(Exception):
    pass


def sheet_name(kind: str, line: str) -> str:
    """Nombre Excel seguro para cédulas estándar y cédulas especiales (ERP/TI, Sistemas, etc.)."""
    prefix = {"SUSCRIPCION": "SUS", "RECLAMO": "REC"}.get(kind, str(kind or "AUD")[:10].upper())
    suffix = str(line or "GENERAL").replace("/", "-").replace("\\", "-").replace(":", "-")
    name = f"{prefix}-{suffix}"[:31]
    return name or "AUDITORIA"


def kind_label(kind: str) -> str:
    return {"SUSCRIPCION": "Suscripción", "RECLAMO": "Reclamo / indemnización"}.get(kind, str(kind or "Auditoría"))


def line_label(line: str) -> str:
    return {"AUTO": "Auto", "PATRIMONIAL": "Patrimonial", "PERSONAL": "Personal"}.get(line, str(line or "General"))


def _canon(obj) -> str:
    return json.dumps(obj, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def _sign(secret: bytes, text: str) -> str:
    return hmac.new(secret, text.encode(), hashlib.sha256).hexdigest()


def _num(v):
    if v is None or v == "":
        return None
    return round(float(v), 2)


def row_hash(results: dict, manual, reason, obs, exc, codes: list[str]) -> str:
    norm = {"r": {c: (results or {}).get(c) or "" for c in codes}, "m": "" if _num(manual) is None else str(_num(manual)),
            "mr": (reason or "").strip(), "o": (obs or "").strip(), "e": (exc or "").strip()}
    return hashlib.sha256(_canon(norm).encode()).hexdigest()[:24]


def _put(ws, row, col, value):
    c = ws.cell(row=row, column=col)
    c.value = value
    if isinstance(value, str) and value[:1] in "=+-@":
        c.data_type = "s"            # texto, nunca fórmula
    return c


# ------------------------------------------------------------------ exportación
def export_workbook(audit: dict, sheets: list[dict], secret: bytes, exported_at: datetime, exported_by: str) -> bytes:
    """audit={id,code,title,unit}; sheets=[{id,kind,line,title,criteria,items:[{id,ref,results,manual_score,manual_reason,observation,info_exception}]}]"""
    Workbook, _, Alignment, Border, Font, PatternFill, Side, get_column_letter, DataValidation = _xlsx()
    wb = Workbook()
    head = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
    fill_head, fill_lock = PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="E7E9EF")
    fill_in = PatternFill("solid", fgColor="FFF9E5")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    base = Font(name=ARIAL, size=10)

    ws0 = wb.active
    ws0.title = "Instrucciones"
    lines = [
        (f"SICAS — Paquete de trabajo sin conexión · {audit['code']}", True),
        (f"{audit['title']} · {audit['unit']}", False),
        (f"Descargado el {exported_at:%d/%m/%Y %H:%M} por {exported_by}", False), ("", False),
        ("CÓMO USAR ESTE ARCHIVO", True),
        ("1. Trabaje en las hojas de las cédulas incluidas en este paquete. Una fila = un registro o control de la muestra.", False),
        ("2. En cada columna de criterio elija: CUMPLE, PARCIAL, NO_CUMPLE, NO_APLICA o EXCEPCION (información no recibida).", False),
        ("3. Si usa EXCEPCION, describa en la última columna qué información no se recibió. Queda como limitación al alcance.", False),
        ("4. La calificación manual (0-100) es opcional y a criterio del Auditor; exige justificación (mínimo 10 caracteres).", False),
        ("5. Puede agregar filas nuevas al final: deje vacía la columna ID y escriba la Referencia.", False),
        ("6. NO modifique las columnas ID ni Referencia de filas existentes, ni las hojas ocultas, ni cambie los nombres de las hojas.", False),
        ("7. Al tener conexión: Auditorías → su auditoría → Sin conexión → Cargar archivo. Verá una vista previa antes de aplicar.", False),
        ("8. Las evidencias (fotos, documentos) se adjuntan en la aplicación al conectarse.", False), ("", False),
        ("Ejemplo de una fila llena:  SIN-000123 | CUMPLE | PARCIAL | NO_CUMPLE | NO_APLICA | … | (manual) 70 | "
         "Subsanado en el sitio | Falta peritaje firmado | EXCEPCION: la sucursal no entregó la denuncia", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws0.cell(row=i, column=1, value=t)
        c.font = Font(name=ARIAL, bold=b, size=12 if i == 1 else 10, color="1F3864" if b else "000000")
    ws0.column_dimensions["A"].width = 120

    crit_ws = wb.create_sheet("Criterios")
    for j, h in enumerate(["Hoja", "Código", "Criterio", "Peso"], 1):
        c = crit_ws.cell(row=1, column=j, value=h); c.font, c.fill = head, fill_head
    crow = 2
    control = {"v": 1, "audit_id": audit["id"], "audit_code": audit["code"], "exported_at": exported_at.isoformat(), "sheets": {}}

    used_names = set(wb.sheetnames)
    for sh in sheets:
        name = sheet_name(sh.get("kind"), sh.get("line"))
        if name in used_names:
            base = name[:27]
            n = 2
            while f"{base}-{n}" in used_names:
                n += 1
            name = f"{base}-{n}"[:31]
        used_names.add(name)
        ws = wb.create_sheet(name)
        codes = [c["code"] for c in sh["criteria"]]
        ws["A1"] = f"{kind_label(sh.get('kind'))} — {line_label(sh.get('line'))}"; ws["A1"].font = Font(name=ARIAL, bold=True, size=12, color="1F3864")
        ws["A2"] = "Criterios: ver hoja «Criterios». Valores: " + " / ".join(RESULTS); ws["A2"].font = Font(name=ARIAL, italic=True, size=9)
        headers = ["ID (no editar)", "Referencia (póliza / siniestro)"] + codes + TAIL
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=HEADER_ROW, column=j, value=h)
            c.font, c.fill, c.border = head, fill_head, box
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[HEADER_ROW].height = 42
        items_ctl = {}
        for r, it in enumerate(sh["items"], FIRST_ROW):
            _put(ws, r, 1, it["id"]).fill = fill_lock
            _put(ws, r, 2, it["ref"]).fill = fill_lock
            for k, code in enumerate(codes):
                _put(ws, r, FIXED_LEFT + 1 + k, (it.get("results") or {}).get(code))
            base_col = FIXED_LEFT + len(codes) + 1
            _put(ws, r, base_col, _num(it.get("manual_score")))
            _put(ws, r, base_col + 1, it.get("manual_reason"))
            _put(ws, r, base_col + 2, it.get("observation"))
            _put(ws, r, base_col + 3, it.get("info_exception"))
            items_ctl[it["id"]] = {"u": str(it.get("updated_at") or ""), "h": row_hash(it.get("results"), it.get("manual_score"), it.get("manual_reason"), it.get("observation"), it.get("info_exception"), codes)}
        last = FIRST_ROW + len(sh["items"]) + 300                  # espacio para filas nuevas
        for r in range(FIRST_ROW, last + 1):
            for j in range(1, len(headers) + 1):
                c = ws.cell(row=r, column=j)
                c.border, c.font = box, base
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if j > FIXED_LEFT:
                    c.fill = fill_in                                 # celdas de captura
                elif r < FIRST_ROW + len(sh["items"]):
                    c.fill = fill_lock                               # ID y referencia existentes: no editar
        dv = DataValidation(type="list", formula1='"' + ",".join(RESULTS) + '"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Valor no válido", error="Use: " + ", ".join(RESULTS))
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(FIXED_LEFT + 1)}{FIRST_ROW}:{get_column_letter(FIXED_LEFT + len(codes))}{last}")
        mcol = get_column_letter(FIXED_LEFT + len(codes) + 1)
        dn = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True, showErrorMessage=True,
                            errorTitle="Calificación", error="Un número entre 0 y 100")
        ws.add_data_validation(dn); dn.add(f"{mcol}{FIRST_ROW}:{mcol}{last}")
        ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 24
        for k in range(len(codes)):
            ws.column_dimensions[get_column_letter(FIXED_LEFT + 1 + k)].width = 12
        for k, w in enumerate([14, 34, 40, 44]):
            ws.column_dimensions[get_column_letter(FIXED_LEFT + len(codes) + 1 + k)].width = w
        ws.freeze_panes = ws.cell(row=FIRST_ROW, column=FIXED_LEFT + 1)
        for c in sh["criteria"]:
            for j, v in enumerate([name, c["code"], c["text"], c["weight"]], 1):
                crit_ws.cell(row=crow, column=j, value=v).font = base
            crow += 1
        control["sheets"][name] = {"id": sh["id"], "kind": sh["kind"], "line": sh["line"], "codes": codes, "items": items_ctl}
    crit_ws.column_dimensions["A"].width = 12; crit_ws.column_dimensions["C"].width = 100

    payload = _canon(control)
    cs = wb.create_sheet("_control")
    cs["A1"], cs["A2"] = MAGIC, _sign(secret, payload)
    for i in range(0, len(payload), CHUNK):
        cs.cell(row=3 + i // CHUNK, column=1, value=payload[i:i + CHUNK]).data_type = "s"
    cs.sheet_state = "hidden"
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


# ------------------------------------------------------------------ importación
def _read_control(wb, secret: bytes) -> dict:
    if "_control" not in wb.sheetnames:
        raise OfflineError("El archivo no es un paquete de SICAS (falta la hoja de control).")
    rows = [r[0] for r in wb["_control"].iter_rows(min_col=1, max_col=1, values_only=True)]
    if len(rows) < 3 or rows[0] != MAGIC:
        raise OfflineError("El archivo no es un paquete de SICAS.")
    payload = "".join(str(x) for x in rows[2:] if x is not None)
    if not hmac.compare_digest(_sign(secret, payload), str(rows[1])):
        raise OfflineError("El archivo fue modificado o no fue generado por este SICAS (la firma de control no coincide).")
    return json.loads(payload)


def parse_workbook(data: bytes, secret: bytes, expected_audit_id: str) -> tuple[dict, list[dict], list[str]]:
    """Devuelve (control, filas, errores). Errores de contenido NO abortan: se informan por fila en la vista previa."""
    if len(data) > MAX_BYTES:
        raise OfflineError("El archivo supera 10 MB.")
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise OfflineError("El archivo no es un libro de Excel (.xlsx).")
    try:
        _, load_workbook, _, _, _, _, _, _, _ = _xlsx()
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise OfflineError("No se pudo leer el archivo de Excel.") from exc
    control = _read_control(wb, secret)
    if control["audit_id"] != expected_audit_id:
        raise OfflineError(f"Este paquete pertenece a otra auditoría ({control['audit_code']}).")
    rows, errors = [], []
    for name, meta in control["sheets"].items():
        if name not in wb.sheetnames:
            errors.append(f"Falta la hoja {name} (¿se cambió el nombre o se eliminó?)."); continue
        codes = meta["codes"]
        ncol = FIXED_LEFT + len(codes) + len(TAIL)
        seen = set()
        for r, vals in enumerate(wb[name].iter_rows(min_row=FIRST_ROW, max_col=ncol, values_only=True), FIRST_ROW):
            if r - FIRST_ROW >= MAX_ROWS:
                errors.append(f"{name}: se excedió el máximo de {MAX_ROWS} filas."); break
            vals = list(vals) + [None] * (ncol - len(vals))
            if all(v in (None, "") for v in vals[1:]) and vals[0] in (None, ""):
                continue
            item_id = str(vals[0]).strip() if vals[0] not in (None, "") else None
            ref = str(vals[1]).strip() if vals[1] not in (None, "") else ""
            row = {"sheet": name, "worksheet_id": meta["id"], "row": r, "item_id": item_id, "ref": ref, "errors": []}
            res = {}
            for k, code in enumerate(codes):
                v = vals[FIXED_LEFT + k]
                v = str(v).strip().upper() if v not in (None, "") else ""
                if v and v not in RESULTS:
                    row["errors"].append(f"{code}: valor «{v}» no válido")
                elif v:
                    res[code] = v
            tail = vals[FIXED_LEFT + len(codes):]
            manual_raw, reason, obs, exc = tail
            try:
                manual = _num(manual_raw)
                if manual is not None and not 0 <= manual <= 100:
                    raise ValueError
            except (TypeError, ValueError):
                manual = None; row["errors"].append("Calificación manual: use un número de 0 a 100")
            reason, obs, exc = [("" if x in (None, "") else str(x).strip()) for x in (reason, obs, exc)]
            if manual is not None and len(reason) < 10:
                row["errors"].append("La calificación manual requiere justificación (mínimo 10 caracteres)")
            if manual is None and reason:
                reason = ""
            if "EXCEPCION" in res.values() and not exc:
                row["errors"].append("Marcó EXCEPCION: describa la información no recibida")
            for label, txt in (("Justificación", reason), ("Observación", obs), ("Excepción", exc)):
                if len(txt) > 2000:
                    row["errors"].append(f"{label}: máximo 2000 caracteres")
            if not item_id and not ref:
                row["errors"].append("Fila nueva sin Referencia")
            if len(ref) > 60:
                row["errors"].append("Referencia: máximo 60 caracteres")
            key = ref.upper()
            if key and key in seen:
                row["errors"].append(f"Referencia repetida en la hoja: {ref}")
            seen.add(key)
            row.update(results=res, manual=manual, reason=reason, obs=obs, exc=exc,
                       hash=row_hash(res, manual, reason, obs, exc, codes))
            rows.append(row)
    wb.close()
    return control, rows, errors


def plan_changes(control: dict, rows: list[dict], current: dict[str, dict], refs_by_sheet: dict[str, set], overwrite: bool = False) -> list[dict]:
    """
    current[item_id] = {"hash": huella del estado ACTUAL en el servidor}; refs_by_sheet[worksheet_id] = referencias existentes.
    Estado por fila: UNCHANGED | UPDATE | NEW | CONFLICT | INVALID.
    """
    plan = []
    for r in rows:
        meta = control["sheets"][r["sheet"]]
        out = {k: r[k] for k in ("sheet", "worksheet_id", "row", "item_id", "ref")}
        if r["errors"]:
            out.update(status="INVALID", detail="; ".join(r["errors"]))
        elif r["item_id"]:
            exported = meta["items"].get(r["item_id"])
            if exported is None:
                out.update(status="INVALID", detail="El ID no pertenece al paquete descargado")
            elif r["hash"] == exported["h"]:
                out.update(status="UNCHANGED", detail="")
            elif r["item_id"] not in current:
                out.update(status="INVALID", detail="El ítem ya no existe en el servidor")
            elif current[r["item_id"]]["hash"] != exported["h"] and not overwrite:
                out.update(status="CONFLICT", detail="Cambió en el servidor después de la descarga")
            else:
                out.update(status="UPDATE", detail="Sobrescribe cambios del servidor" if current[r["item_id"]]["hash"] != exported["h"] else "")
        else:
            if r["ref"].upper() in {x.upper() for x in refs_by_sheet.get(r["worksheet_id"], set())}:
                out.update(status="CONFLICT", detail="Ya existe un ítem con esa referencia")
            else:
                out.update(status="NEW", detail="")
        out["data"] = {k: r[k] for k in ("results", "manual", "reason", "obs", "exc")}
        plan.append(out)
    return plan


def summarize(plan: list[dict]) -> dict:
    s = {k: 0 for k in ("UPDATE", "NEW", "UNCHANGED", "CONFLICT", "INVALID")}
    for p in plan:
        s[p["status"]] += 1
    return s
